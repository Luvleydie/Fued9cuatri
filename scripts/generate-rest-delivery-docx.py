"""Build the separate REST delivery report from verified local evidence.

Usage: python generate-rest-delivery-docx.py --commit <published code SHA>
Requires python-docx and Pillow. The --commit is the published implementation
commit, before the separate documentation commit that adds this report.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
REPO = "https://github.com/Luvleydie/Fued9cuatri"
DEFAULT_BRANCH = "entrega/api-rest-2026-10-14"


def hyperlink(paragraph, label, url):
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "15517B")
    prop.append(color)
    run.append(prop)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def para(doc, text, style=None):
    paragraph = doc.add_paragraph(text, style)
    for run in paragraph.runs:
        run.font.bold = style == "Title"
        run.font.italic = False
    return paragraph


def table(doc, headers, rows, widths):
    item = doc.add_table(rows=1, cols=len(headers))
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.autofit = False
    for column, width in zip(item.columns, widths):
        column.width = Inches(width)
    for index, cell in enumerate(item.rows[0].cells):
        cell.text = headers[index]
        cell.width = Inches(widths[index])
    for values in rows:
        row = item.add_row()
        for index, cell in enumerate(row.cells):
            cell.text = str(values[index])
            cell.width = Inches(widths[index])
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        edge = OxmlElement(f"w:{name}")
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    item._tbl.tblPr.append(borders)
    for row_index, row in enumerate(item.rows):
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        props.append(no_split)
        if row_index == 0:
            props.append(OxmlElement("w:tblHeader"))
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for name in ["top", "left", "bottom", "right"]:
                margin = OxmlElement(f"w:{name}")
                margin.set(qn("w:w"), "95")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "24475B" if row_index == 0 else ("F3F6F8" if row_index % 2 else "FFFFFF"))
            properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9.3)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "000000")
    para(doc, "").paragraph_format.space_after = Pt(0)
    return item


def figure(doc, filename, caption, max_height=3.2, crop_bottom=0):
    path = ROOT / "docs/screenshots-rest" / filename
    with Image.open(path) as picture:
        width, height = picture.size
    height *= 1 - crop_bottom
    physical_width = min(6.9, max_height * width / height)
    paragraph = para(doc, "")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(physical_width), height=Inches(physical_width * height / width))
    if crop_bottom:
        # Recorte de presentación en Word; el PNG de evidencia queda intacto.
        rectangle = OxmlElement("a:srcRect")
        rectangle.set("b", str(round(crop_bottom * 100000)))
        inline._inline.xpath(".//pic:blipFill")[0].insert(1, rectangle)
    inline._inline.docPr.set("descr", caption)
    line = para(doc, caption, "Caption")
    line.paragraph_format.space_after = Pt(12)


def heading(doc, text, new_page=False):
    if new_page:
        doc.add_page_break()
    return doc.add_heading(text, level=1)


def code(doc, text):
    p = para(doc, "")
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(text.splitlines()):
        run = p.add_run(("\n" if index else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.bold = False
    return p


def build(args):
    if not re.fullmatch(r"[0-9a-f]{40}", args.commit):
        raise ValueError("--commit must be the full SHA of the verified published code commit")
    evidence_path = ROOT / "docs/evidencia/rest-verificacion.json"
    evidence = json.loads(evidence_path.read_text(encoding="utf-8-sig"))
    if not evidence.get("passed") or not evidence.get("cleanup", {}).get("completed"):
        raise ValueError("Browser verification or cleanup was unsuccessful")
    if len(evidence["checks"]) != 12 or not all(check.get("passed") for check in evidence["checks"]):
        raise ValueError("Expected twelve passing browser checks")
    quality = json.loads((ROOT / "docs/evidencia/rest-calidad.json").read_text(encoding="utf-8-sig"))
    if not all(check.get("passed") for check in quality["checks"]):
        raise ValueError("Quality evidence contains an unsuccessful check")
    screenshots = ["01-loading-get.png", "02-get-usuarios.png", "03-loading-post.png", "04-post-creado.png",
                   "05-put-actualizado.png", "06-delete-eliminado.png", "07-error-controlado.png", "08-recuperacion.png"]
    for filename in screenshots:
        if not (ROOT / "docs/screenshots-rest" / filename).is_file():
            raise FileNotFoundError(filename)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    section.footer_distance = Inches(.27)
    styles = doc.styles
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].font.bold = name in ["Title", "Heading 1", "Heading 2"]
        styles[name].font.italic = False
        for border in styles[name].element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.05
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(7)
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].paragraph_format.space_after = Pt(15)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].paragraph_format.space_before = Pt(9)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].paragraph_format.space_before = Pt(7)
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = False
    for name in ["Title", "Subtitle", "Heading 1", "Heading 2"]:
        styles[name].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("NovaCart  |  ")
    run.font.size = Pt(8)
    run.font.bold = False
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Conexión de NovaCart con una API REST"
    doc.core_properties.subject = "Evidencias de integración Ionic Angular PHP MySQL"
    doc.core_properties.author = "Proyecto NovaCart"

    para(doc, "Conexión de NovaCart\ncon una API REST", "Title")
    para(doc, "Entrega de implementación y evidencias de funcionamiento", "Subtitle")
    para(doc, "14 de octubre de 2026")
    para(doc, "NovaCart consulta y administra usuarios desde una aplicación Ionic con Angular. Las solicitudes HTTP llegan a una API PHP, que valida los datos y conserva los cambios en MySQL. Esta entrega reúne el código, los resultados de comprobación y las capturas del flujo de consulta, alta, actualización y eliminación.")
    para(doc, "La actualización se implementa con PUT, una de las dos opciones solicitadas en PUT/PATCH. Los servicios Angular separan las peticiones de la interfaz. Durante las solicitudes se muestran indicadores de carga; los errores se presentan al usuario sin perder los datos del formulario.")
    heading(doc, "Cobertura de los requisitos")
    table(doc, ["Requisito", "Implementación", "Evidencia"], [
        ["GET", "Consulta de usuarios con UsersService.getUsers", "GET /users · 200"],
        ["POST", "Alta de usuario con UsersService.createUser", "POST /users · 201"],
        ["PUT o PATCH", "Edición con UsersService.updateUser mediante PUT", "PUT /users/{id} · 200"],
        ["DELETE", "Baja con UsersService.deleteUser y confirmación", "DELETE /users/{id} · 204"],
        ["Servicios Angular", "UsersService y CartService inyectables", "Código en src/app/services"],
        ["Respuestas y errores", "Validación de datos, try/catch/finally y mensajes", "409 real, 503 simulado y recuperación"],
        ["Loading", "ion-spinner, isLoading, isSubmitting y bloqueo de controles", "Capturas durante GET y POST"],
        ["Git y capturas", "Rama de entrega y registro verificable de pruebas", "Repositorio y evidencias adjuntas"],
    ], [1.15, 3.7, 2.05])
    para(doc, "La demostración utiliza la aplicación web instalada en http://localhost/novacart/ y la API local de XAMPP. Las operaciones CRUD se ejecutan contra PHP y MySQL; el error 503 se provoca de forma controlada en el navegador para comprobar la interfaz.")

    heading(doc, "Servicios Angular y flujo de solicitudes", True)
    para(doc, "UsersService está registrado con @Injectable({ providedIn: 'root' }) y se inyecta en HomePage. Sus métodos devuelven promesas tipadas. El cliente Axios compartido usa la URL del entorno, envía JSON, agrega el token Bearer y limita cada petición a 15 segundos.")
    table(doc, ["Archivo", "Responsabilidad"], [
        ["src/app/services/users.service.ts", "GET, POST, PUT y DELETE de usuarios; validación y actualización de caché."],
        ["src/app/services/cart.service.ts", "Consulta de catálogo y carrito, cambio de cantidades y eliminación de artículos."],
        ["src/app/core/api/axios-client.ts", "Cliente HTTP común, encabezados, token y estado de conexión."],
        ["src/app/core/api/api-errors.ts", "Mensajes de red, timeout, sesión, validación y fallos HTTP."],
        ["src/app/pages/home/home.page.ts", "Invocación del servicio, actualización de lista y estados de la interfaz."],
        ["server/endpoints/users.php", "Rutas de usuarios, validación, SQL y respuestas HTTP."],
    ], [3.1, 3.8])
    doc.add_heading("Peticiones del servicio", level=2)
    para(doc, "Estas son las llamadas HTTP centrales de UsersService. En el archivo completo están rodeadas por controles de sesión, validación de respuestas e invalidación de caché.")
    code(doc, "api.get<UsersResponse>('/users');\napi.post<User>('/users', input);\napi.put<User>(`/users/${id}`, input);\napi.delete(`/users/${id}`);")
    doc.add_heading("Respuestas errores y carga", level=2)
    para(doc, "La lista se modifica cuando la API confirma la operación. GET entrega { users: [...] }; POST y PUT devuelven el usuario guardado. DELETE devuelve 204 sin cuerpo. Las respuestas se proyectan a los modelos esperados antes de usarlas.")
    para(doc, "HomePage activa isLoading al consultar e isSubmitting al guardar o eliminar. Los controles se bloquean para evitar solicitudes repetidas. Cada bloque finally restablece el estado, incluso si ocurre un error; los spinners Ionic hacen visible la espera.")
    para(doc, "El catch traduce el error y conserva el formulario. Ante un fallo incierto de escritura se pide actualizar antes de repetir para evitar duplicados. Un 401 limpia la sesión y vuelve al acceso. Los errores de validación se asocian a los campos cuando la API los proporciona.")

    heading(doc, "Consulta GET y estado de carga", True)
    para(doc, "Al pulsar Actualizar se solicita GET /users. Para capturar la espera se retuvo temporalmente la solicitud antes de enviarla a la API real. La prueba comprueba el indicador de carga y el bloqueo de acciones mientras la solicitud está pendiente.")
    figure(doc, "01-loading-get.png", "Figura 1. Consulta pendiente con indicador de carga y controles bloqueados.")
    figure(doc, "02-get-usuarios.png", "Figura 2. GET finalizado con HTTP 200 y usuarios recibidos de la API.")

    heading(doc, "Creación POST de un usuario", True)
    para(doc, "El formulario envía los datos del usuario a POST /users. La respuesta real 201 confirma el alta y el usuario aparece en el directorio. La prueba también verifica que el formulario queda bloqueado durante el envío. La contraseña no se incluye en las evidencias.")
    figure(doc, "03-loading-post.png", "Figura 3. Detalle del POST pendiente con Procesando solicitud y Actualizar deshabilitado.", crop_bottom=.55)
    figure(doc, "04-post-creado.png", "Figura 4. Usuario creado en el directorio tras la respuesta HTTP 201.")

    heading(doc, "Actualización PUT y eliminación DELETE", True)
    para(doc, "Editar carga los datos en el formulario y Guardar cambios envía PUT /users/{id}. El nombre actualizado se comprueba después de recargar. Eliminar solicita confirmación y envía DELETE al mismo identificador; tras el 204, la fila desaparece y la baja persiste.")
    figure(doc, "05-put-actualizado.png", "Figura 5. Actualización confirmada por PUT con respuesta HTTP 200.")
    figure(doc, "06-delete-eliminado.png", "Figura 6. Eliminación confirmada por DELETE con respuesta HTTP 204.")

    heading(doc, "Manejo de error y recuperación", True)
    para(doc, "Un segundo intento de crear el mismo usuario devuelve HTTP 409 real. La aplicación informa que el nombre ya está registrado y conserva el formulario; no se crea un duplicado. La captura adicional se conserva como 09-error-409-real.png en el repositorio.")
    para(doc, "La prueba intercepta únicamente un POST de prueba y devuelve HTTP 503 para simular un servicio no disponible. Esta captura comprueba el manejo del fallo en la interfaz; no representa una caída real de MySQL. Los datos del formulario se conservan y se solicita actualizar antes de repetir.")
    figure(doc, "07-error-controlado.png", "Figura 7. Mensaje del error 503 simulado que solicita actualizar antes de repetir.")
    figure(doc, "08-recuperacion.png", "Figura 8. Datos conservados en el formulario después de recuperar la consulta.")

    heading(doc, "Código publicado y comprobación", True)
    para(doc, "El código de implementación y las evidencias de las pruebas están publicados en GitHub. El commit siguiente identifica la versión de código verificada; el documento se incorpora en un commit posterior de la misma rama.")
    p = para(doc, "Repositorio: ")
    hyperlink(p, "Luvleydie/Fued9cuatri", REPO)
    p = para(doc, "Rama de entrega: ")
    hyperlink(p, args.branch, f"{REPO}/tree/{args.branch}")
    p = para(doc, "Commit de código: ")
    hyperlink(p, args.commit, f"{REPO}/commit/{args.commit}")
    table(doc, ["Comprobación", "Resultado"], [
        ["npm test -- --watch=false", "322 pruebas aprobadas en 22 archivos"],
        ["npm run test:api", "35 comprobaciones aprobadas con MySQL"],
        ["npm run lint", "Sin errores"],
        ["npm run deploy:xampp", "Compilación y despliegue local completados"],
        ["npm run test:rest-delivery", "12 comprobaciones aprobadas y 0 usuarios temporales restantes"],
    ], [3.4, 3.5])
    doc.add_heading("Cómo reproducir la demostración", level=2)
    para(doc, "Con Node y XAMPP instalados, activa Apache y MySQL y prepara la base novacart. Desde una copia del repositorio ejecuta:")
    code(doc, f"git switch {args.branch}\nnpm ci\nnpm run db:setup\nnpm run deploy:xampp\nnpm run test:rest-delivery")
    para(doc, "Abre http://localhost/novacart/ e inicia sesión con la cuenta de demostración indicada en README.md. En Inicio, actualiza el directorio, crea un usuario temporal, edítalo y elimínalo. Las pruebas automatizadas limpian sus cuentas temporales al terminar.")
    doc.add_heading("Archivos de evidencia", level=2)
    para(doc, "docs/screenshots-rest contiene las capturas. docs/evidencia/rest-verificacion.json conserva los resultados de la interfaz y el registro de las peticiones. docs/evidencia/rest-calidad.json registra las pruebas de código, API, lint y compilación. scripts/verify-rest-delivery.mjs reproduce la verificación.")
    para(doc, "El alcance comprobado es la integración web Ionic y Angular con PHP y MySQL. Esta entrega no documenta una ejecución del APK en hardware Android.")
    out = ROOT / "output/NovaCart_Entrega_API_REST.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(json.dumps({"output": str(out), "codeCommit": args.commit, "evidence": str(evidence_path), "screenshots": len(screenshots)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--commit", required=True)
    parser.add_argument("--branch", default=DEFAULT_BRANCH)
    build(parser.parse_args())
